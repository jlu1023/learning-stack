"""Convert a source file into text, plus page images where the text is not enough.

`/learn:ingest` and `/learn:connect` read the result: the text first, and a
page image wherever the PDF's text layer is missing (a scan) or garbles the
mathematics. Only free, standard converters are used, each optional:

| Source | Converter |
|---|---|
| PDF | pypdfium2 (with Pillow for page images), else poppler's pdftotext / pdftoppm |
| DOCX | pandoc (Word equations come out as LaTeX) |
| PPTX | python-pptx for slide text and speaker notes; LibreOffice for slide images |
| HTML | pandoc |
| Markdown, text | read as they are |
| Images | none: Claude reads them directly |

A missing converter raises MissingTool, whose message says how to install it.
Output goes to one folder: `text.md`, `page-NNN.png` and `manifest.json`.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

PDF = {".pdf"}
DOCX = {".docx"}
PPTX = {".pptx"}
HTML = {".html", ".htm"}
TEXT = {".md", ".markdown", ".txt"}
IMAGE = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
SUPPORTED = PDF | DOCX | PPTX | HTML | TEXT | IMAGE

INSTALL = {
    "pypdfium2": "python -m pip install pypdfium2 pillow",
    "pillow": "python -m pip install pillow",
    "poppler": "macOS: brew install poppler; Debian or Ubuntu: sudo apt install poppler-utils; "
               "Windows: install a Poppler build and put its bin folder on PATH",
    "pandoc": "Windows: winget install --id JohnMacFarlane.Pandoc; macOS: brew install pandoc; "
              "Debian or Ubuntu: sudo apt install pandoc; or see https://pandoc.org/installing.html",
    "python-pptx": "python -m pip install python-pptx",
    "libreoffice": "Windows: winget install TheDocumentFoundation.LibreOffice; macOS: brew install --cask "
                   "libreoffice; Debian or Ubuntu: sudo apt install libreoffice-impress",
}

# characters a text layer produces for mathematics it cannot represent well
MATHS_CHARS = set("∑∫∮√∞∂∇≤≥≠≈≡∈∉⊂⊆⊃⊇∪∩∀∃→⇒⇔↦±×÷∘⋅′″αβγδεζηθικλμνξπρστυφχψωΓΔΘΛΞΠΣΦΨΩ")
SCAN_CHARS = 40          # fewer characters than this on a page: treat it as a scan


class MissingTool(RuntimeError):
    def __init__(self, tool: str, purpose: str):
        self.tool, self.install = tool, INSTALL[tool]
        if tool == "pypdfium2":
            self.install += f" (or poppler's pdftotext instead: {INSTALL['poppler']})"
        super().__init__(f"{purpose} needs {tool}. Install: {self.install}")


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def default_out(root: Path, src: Path) -> Path:
    stem = re.sub(r"[^\w.-]+", "-", src.stem).strip("-")[:60] or "source"
    return root / ".converted" / f"{stem}-{file_hash(src)[:8]}"


def render_reason(text: str) -> str | None:
    """Why a page should also be read as an image, or None if its text will do."""
    t = text.strip()
    if len(t) < SCAN_CHARS:
        return "little or no text layer (a scan or a picture)"
    odd = sum(1 for c in t if c in MATHS_CHARS or c == "�" or "" <= c <= "")
    lines = [l for l in t.splitlines() if l.strip()]
    short = sum(1 for l in lines if len(l.strip()) <= 3)
    if odd >= 3 or "�" in t:
        return "mathematics in the text layer, which loses layout (fractions, sub- and superscripts)"
    if lines and short / len(lines) > 0.3:
        return "text broken into fragments, typical of typeset equations"
    return None


def _which(*names: str) -> str | None:
    return next((p for p in (shutil.which(n) for n in names) if p), None)


# ---------- PDF ----------
def _pdf_pypdfium2(src: Path, out: Path, render: str, dpi: int) -> dict:
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(src))
    pages, texts, warnings = [], [], []
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            tp = page.get_textpage()
            text = tp.get_text_range().replace("\r\n", "\n")
            tp.close()
            why = "all pages rendered (--render all)" if render == "all" else (
                render_reason(text) if render == "auto" else None)
            img = None
            if why:
                try:
                    pil = page.render(scale=dpi / 72).to_pil()
                except (ImportError, AttributeError):
                    warnings.append(f"page images need Pillow: {INSTALL['pillow']}")
                    render, why = "none", None
                else:
                    img = f"page-{i + 1:03d}.png"
                    pil.save(out / img)
            texts.append(text)
            pages.append({"page": i + 1, "chars": len(text.strip()), "image": img, "why": why})
            page.close()
    finally:
        pdf.close()
    return {"backend": "pypdfium2", "pages": pages, "texts": texts, "warnings": warnings}


def _pdf_poppler(src: Path, out: Path, render: str, dpi: int) -> dict:
    exe = _which("pdftotext")
    if not exe:
        raise MissingTool("pypdfium2", "Reading a PDF")
    proc = subprocess.run([exe, "-layout", "-enc", "UTF-8", str(src), "-"], capture_output=True)
    if proc.returncode != 0:
        raise RuntimeError(f"pdftotext failed: {proc.stderr.decode('utf-8', 'replace').strip()[:300]}")
    texts = proc.stdout.decode("utf-8", "replace").split("\f")
    if texts and not texts[-1].strip():
        texts = texts[:-1]
    ppm, pages, warnings = _which("pdftoppm"), [], []
    if render != "none" and not ppm:
        warnings.append(f"page images need pdftoppm (poppler) or pypdfium2: {INSTALL['pypdfium2']}")
    for i, text in enumerate(texts):
        why = "all pages rendered (--render all)" if render == "all" else (
            render_reason(text) if render == "auto" else None)
        img = None
        if why and ppm:
            base = out / f"page-{i + 1:03d}"
            subprocess.run([ppm, "-r", str(dpi), "-f", str(i + 1), "-l", str(i + 1), "-png",
                            "-singlefile", str(src), str(base)], capture_output=True)
            if base.with_suffix(".png").exists():
                img = base.with_suffix(".png").name
        pages.append({"page": i + 1, "chars": len(text.strip()), "image": img, "why": why})
    return {"backend": "poppler (pdftotext)", "pages": pages, "texts": texts, "warnings": warnings}


def _pdf(src: Path, out: Path, render: str, dpi: int) -> dict:
    try:
        import pypdfium2  # noqa: F401
    except ImportError:
        return _pdf_poppler(src, out, render, dpi)
    return _pdf_pypdfium2(src, out, render, dpi)


# ---------- other formats ----------
def _pandoc(src: Path, out: Path, fmt: str) -> dict:
    exe = _which("pandoc")
    if not exe:
        raise MissingTool("pandoc", f"Reading a {fmt.upper()} file")
    proc = subprocess.run([exe, str(src), "-f", fmt, "-t", "markdown-raw_html-native_divs-native_spans",
                           "--wrap=none", f"--extract-media={out / 'media'}", "-o", str(out / "text.md")],
                          capture_output=True)
    if proc.returncode != 0:
        raise RuntimeError(f"pandoc failed: {proc.stderr.decode('utf-8', 'replace').strip()[:300]}")
    media = sorted(p.relative_to(out).as_posix() for p in (out / "media").rglob("*") if p.is_file()) \
        if (out / "media").exists() else []
    return {"backend": "pandoc", "media": media, "warnings": []}


def _pptx(src: Path, out: Path, render: str, dpi: int) -> dict:
    try:
        from pptx import Presentation
    except ImportError:
        raise MissingTool("python-pptx", "Reading a PPTX file") from None
    prs, L = Presentation(str(src)), []
    for i, slide in enumerate(prs.slides, 1):
        L.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                L += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                L += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            L += ["", "Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip()]
        L.append("")
    (out / "text.md").write_text("\n".join(L), encoding="utf-8")
    res = {"backend": "python-pptx", "slides": len(prs.slides), "pages": [], "warnings": []}
    if render == "none":
        return res
    office = _which("soffice", "libreoffice")
    if not office:
        res["warnings"].append("slide text only: rendering the slides as images needs LibreOffice "
                               f"({INSTALL['libreoffice']})")
        return res
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                       capture_output=True, timeout=300)
        pdf = Path(tmp) / f"{src.stem}.pdf"
        if pdf.exists():
            r = _pdf(pdf, out, "all", dpi)
            res["pages"], res["warnings"] = r["pages"], res["warnings"] + r["warnings"]
        else:
            res["warnings"].append("LibreOffice could not convert the slides to PDF")
    return res


# ---------- entry point ----------
def convert(src: Path, out: Path, render: str = "auto", dpi: int = 150) -> dict:
    """Convert `src` into `out`. Returns the manifest (also saved as manifest.json)."""
    src = Path(src)
    if not src.is_file():
        raise FileNotFoundError(f"no such file: {src}")
    ext = src.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"unsupported file type '{ext}'; supported: {', '.join(sorted(SUPPORTED))}")
    m = {"source": str(src), "hash": file_hash(src), "out": str(out), "kind": ext[1:]}
    if ext in IMAGE:
        # nothing to convert: Claude reads the image itself
        return {**m, "backend": "none", "images": [str(src)], "text_file": None, "warnings": [],
                "note": "read the image directly; mark any word you cannot read as [?: ...]"}
    if ext in TEXT:
        return {**m, "backend": "none", "text_file": str(src), "warnings": []}
    out.mkdir(parents=True, exist_ok=True)
    if ext in PDF:
        r = _pdf(src, out, render, dpi)
        texts = r.pop("texts")
        (out / "text.md").write_text("".join(f"\n\n<!-- page {i} -->\n\n{t.strip()}"
                                             for i, t in enumerate(texts, 1)).lstrip(), encoding="utf-8")
    elif ext in DOCX:
        r = _pandoc(src, out, "docx")
    elif ext in HTML:
        r = _pandoc(src, out, "html")
    else:
        r = _pptx(src, out, render, dpi)
    m.update(r)
    m["text_file"] = str(out / "text.md")
    m["images"] = [str(out / p["image"]) for p in m.get("pages", []) if p.get("image")]
    (out / "manifest.json").write_text(json.dumps(m, indent=1), encoding="utf-8")
    return m
